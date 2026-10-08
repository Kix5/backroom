"""Local document index: read files, chunk them, embed with the local model, search.

Everything is stored as plain files in the index folder:
  chunks.jsonl   one JSON line per chunk (source file, location, text)
  vectors.npy    normalised embedding for each chunk
  files.json     size + modified time of each indexed file (for incremental updates)
"""
from __future__ import annotations

import io
import json
import os
import re
import zipfile
from pathlib import Path

import numpy as np

TEXT_EXT = {".md", ".txt", ".mmd", ".csv", ".json", ".tex", ".yaml", ".yml"}
DOC_EXT = {".pdf", ".docx", ".pptx", ".potx"} | TEXT_EXT


# ------------------------------------------------------------------ readers --
def _open_pptx(path):
    from pptx import Presentation
    if str(path).lower().endswith(".potx"):  # python-pptx refuses templates; relabel in memory
        buf = io.BytesIO()
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(b"presentationml.template.main+xml",
                                        b"presentationml.presentation.main+xml")
                zout.writestr(item, data)
        buf.seek(0)
        return Presentation(buf)
    return Presentation(path)


def read_file(path: Path) -> list[tuple[str, str]]:
    """Return [(location, text)] sections for one file."""
    ext = path.suffix.lower()
    if ext in TEXT_EXT:
        return [("", path.read_text(encoding="utf-8", errors="ignore"))]
    if ext == ".pdf":
        from pypdf import PdfReader
        return [(f"page {i}", p.extract_text() or "")
                for i, p in enumerate(PdfReader(str(path)).pages, 1)]
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return [("", "\n".join(parts))]
    if ext in (".pptx", ".potx"):
        prs, out = _open_pptx(path), []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shp in slide.shapes:
                if shp.has_text_frame:
                    texts.append(shp.text_frame.text)
                if getattr(shp, "has_table", False) and shp.has_table:
                    for row in shp.table.rows:
                        texts.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide:
                texts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
            out.append((f"slide {i}", "\n".join(t for t in texts if t.strip())))
        return out
    return []


def chunk(text: str, size: int, overlap: int) -> list[str]:
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= size:
        return [text] if text else []
    out, start = [], 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):  # break on a paragraph or sentence boundary if possible
            cut = max(text.rfind("\n\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if cut > start:
                end = cut + 1
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in out if c]


# -------------------------------------------------------------------- index --
class Index:
    def __init__(self, folder: str, llm, rag_cfg: dict):
        self.dir = Path(folder)
        self.llm = llm
        self.cfg = rag_cfg
        self.chunks: list[dict] = []
        self.vectors = np.zeros((0, 0), dtype=np.float32)
        self.files: dict = {}
        if (self.dir / "chunks.jsonl").exists():
            with open(self.dir / "chunks.jsonl", encoding="utf-8") as f:
                self.chunks = [json.loads(line) for line in f if line.strip()]
            self.vectors = np.load(self.dir / "vectors.npy")
            self.files = json.loads((self.dir / "files.json").read_text(encoding="utf-8"))

    def save(self):
        self.dir.mkdir(parents=True, exist_ok=True)
        with open(self.dir / "chunks.jsonl", "w", encoding="utf-8") as f:
            f.writelines(json.dumps(c, ensure_ascii=False) + "\n" for c in self.chunks)
        np.save(self.dir / "vectors.npy", self.vectors)
        (self.dir / "files.json").write_text(json.dumps(self.files, indent=1), encoding="utf-8")

    def _embed(self, texts, prefix):
        v = np.asarray(self.llm.embed([prefix + t for t in texts]), dtype=np.float32)
        return v / np.maximum(np.linalg.norm(v, axis=1, keepdims=True), 1e-9)

    def add_paths(self, paths: list[str], log=print) -> tuple[int, int]:
        files = []
        for p in paths:
            p = Path(p)
            if p.is_dir():
                files += [f for f in sorted(p.rglob("*")) if f.suffix.lower() in DOC_EXT
                          and not any(part.startswith(".") for part in f.parts)]
            elif p.suffix.lower() in DOC_EXT:
                files.append(p)
            else:
                log(f"  skip (unsupported type): {p}")
        added = skipped = 0
        for f in files:
            key = str(f.resolve())
            st = os.stat(f)
            sig = {"size": st.st_size, "mtime": int(st.st_mtime)}
            if self.files.get(key) == sig:
                skipped += 1
                continue
            try:
                sections = read_file(f)
            except Exception as e:  # unreadable file: report and continue
                log(f"  ! could not read {f.name}: {e}")
                continue
            pieces = []
            for loc, text in sections:
                for c in chunk(text, self.cfg.get("chunk_chars", 1200), self.cfg.get("overlap_chars", 200)):
                    pieces.append({"source": key, "name": f.name, "loc": loc, "text": c})
            keep = [i for i, c in enumerate(self.chunks) if c["source"] != key]
            self.chunks = [self.chunks[i] for i in keep]
            self.vectors = self.vectors[keep] if len(self.vectors) else self.vectors
            if pieces:
                vec = self._embed([p["text"] for p in pieces], self.cfg.get("doc_prefix", ""))
                self.vectors = vec if not len(self.vectors) else np.vstack([self.vectors, vec])
                self.chunks += pieces
            self.files[key] = sig
            added += 1
            log(f"  + {f.name}: {len(pieces)} chunks")
        self.save()
        return added, skipped

    def search(self, query: str, k: int | None = None, only: str | None = None) -> list[dict]:
        if not self.chunks:
            return []
        k = k or self.cfg.get("top_k", 6)
        q = self._embed([query], self.cfg.get("query_prefix", ""))[0]
        scores = self.vectors @ q
        order = np.argsort(-scores)
        out = []
        for i in order:
            c = self.chunks[int(i)]
            if only and only.lower() not in c["name"].lower():
                continue
            out.append({**c, "score": float(scores[int(i)])})
            if len(out) >= k:
                break
        return out


def format_context(hits: list[dict], limit_chars: int = 8000) -> str:
    parts, used = [], 0
    for i, h in enumerate(hits, 1):
        loc = f", {h['loc']}" if h["loc"] else ""
        block = f"[{i}] ({h['name']}{loc})\n{h['text']}\n"
        if used + len(block) > limit_chars:
            break
        parts.append(block)
        used += len(block)
    return "\n".join(parts)
