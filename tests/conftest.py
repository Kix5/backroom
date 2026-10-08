"""Shared fixtures: a fake OpenAI-compatible server standing in for Ollama / LM Studio."""
from __future__ import annotations

import hashlib
import json
import re
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest

GOOD_DIAGRAM = '''<think>planning the layout</think>
```mermaid
flowchart LR
  subgraph clients["Clients"]
    web["Web app"]
  end
  api["API<br/>REST"]
  db[("Database")]
  web --> api --> db
  class web edge
  class api app
  class db data
```'''

BROKEN_DIAGRAM = '```mermaid\nflowchart LR\n  a["A"] --> b["B\n```'


def fake_embedding(text: str, dim: int = 32) -> list[float]:
    v = [0.0] * dim
    for w in re.findall(r"\w+", text.lower()):
        v[int(hashlib.md5(w.encode()).hexdigest(), 16) % dim] += 1.0
    return v


class FakeServer:
    def __init__(self):
        self.requests: list[dict] = []
        self.diagram_replies: list[str] = [GOOD_DIAGRAM]
        self.text_reply = "Answer from documents [1]."
        server = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                server.requests.append({"path": self.path, "body": body})
                if self.path.endswith("/embeddings"):
                    inp = body["input"] if isinstance(body["input"], list) else [body["input"]]
                    out = {"data": [{"index": i, "embedding": fake_embedding(t)} for i, t in enumerate(inp)]}
                else:
                    system = body["messages"][0]["content"]
                    if "Mermaid" in system:
                        reply = server.diagram_replies.pop(0) if len(server.diagram_replies) > 1 \
                            else server.diagram_replies[0]
                    else:
                        reply = server.text_reply
                    out = {"choices": [{"message": {"role": "assistant", "content": reply}}]}
                data = json.dumps(out).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

        self.httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self.httpd.server_address[1]}/v1"
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def chats(self):
        return [r for r in self.requests if r["path"].endswith("/chat/completions")]


@pytest.fixture
def fake_server():
    s = FakeServer()
    yield s
    s.httpd.shutdown()


@pytest.fixture
def workdir(tmp_path, fake_server, monkeypatch):
    """A project folder with backroom.json pointing at the fake server."""
    cfg = {
        "backend": "ollama",
        "backends": {
            "ollama": {"base_url": fake_server.url, "chat_model": "test-chat", "embed_model": "test-embed"},
            "lmstudio": {"base_url": fake_server.url, "chat_model": "test-chat", "embed_model": "test-embed"},
        },
    }
    (tmp_path / "backroom.json").write_text(json.dumps(cfg), encoding="utf-8")
    monkeypatch.chdir(tmp_path)
    return tmp_path


@pytest.fixture
def template(tmp_path):
    """A small PowerPoint file with the default Office theme and one text slide."""
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Diagram style rules"
    slide.placeholders[1].text = "Firewalls use the hexagon shape. Use sentence case."
    path = tmp_path / "template.pptx"
    prs.save(path)
    return path
